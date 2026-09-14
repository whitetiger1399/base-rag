"""Structure-aware ingestion for the Malawi IDSR Excel booklets."""

import json
import re
import unicodedata
from pathlib import Path
from typing import Iterable, Iterator, List, Sequence, Tuple

from openpyxl import load_workbook

from .models import Chunk


HEADING_RE = re.compile(
    r"^(?:SECTION|CHAPTER|PART|ANNEX|APPENDIX)\b|^\d+(?:\.\d+){1,4}\s+",
    re.IGNORECASE,
)
SPACE_RE = re.compile(r"\s+")
SENTENCE_BOUNDARY_RE = re.compile(r"(?<=[.!?;:])\s+(?=[A-Z0-9])")


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("\ufdd0", " ").replace("\ufffc", " ")
    return SPACE_RE.sub(" ", text).strip()


def normalize_paragraph_number(value: object, fallback: int) -> int:
    if value is None:
        return fallback
    try:
        return int(float(str(value)))
    except (TypeError, ValueError):
        match = re.search(r"\d+", str(value))
        return int(match.group()) if match else fallback


def is_heading(text: str) -> bool:
    """Identify short structural labels while avoiding sentence-like paragraphs."""
    if not text or len(text) > 160:
        return False
    words = text.split()
    if len(words) > 24:
        return False
    if HEADING_RE.match(text):
        return not text.endswith(('.', '?', '!')) or len(words) <= 12
    alpha = [c for c in text if c.isalpha()]
    return bool(alpha) and text.upper() == text and len(words) >= 2


def clean_heading(text: str) -> str:
    # TOC headings often end in a page number; removing it gives cleaner metadata.
    return re.sub(r"\s+\d+$", "", text).strip(" :-") or "Document introduction"


def split_long_text(text: str, max_chars: int) -> List[str]:
    """Split an oversized source paragraph without losing any source text."""
    if len(text) <= max_chars:
        return [text]
    sentences = SENTENCE_BOUNDARY_RE.split(text)
    parts: List[str] = []
    current = ""
    for sentence in sentences:
        if len(sentence) > max_chars:
            if current:
                parts.append(current)
                current = ""
            for start in range(0, len(sentence), max_chars):
                parts.append(sentence[start : start + max_chars].strip())
            continue
        projected = f"{current} {sentence}".strip()
        if current and len(projected) > max_chars:
            parts.append(current)
            current = sentence
        else:
            current = projected
    if current:
        parts.append(current)
    return [part for part in parts if part]


def load_booklet_rows(path: Path) -> List[Tuple[int, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows: List[Tuple[int, str]] = []
        for position, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            paragraph = values[0] if values else position
            text = normalize_text(values[1] if len(values) > 1 else None)
            if text:
                rows.append((normalize_paragraph_number(paragraph, position), text))
        return rows
    finally:
        workbook.close()


def _doc_id(path: Path) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", path.stem).strip("_")


def _topic(section: str) -> str:
    topic = re.sub(r"^\d+(?:\.\d+)*\s*", "", section)
    return topic[:160].strip(" :-")


def _emit_chunk(
    path: Path,
    section: str,
    rows: Sequence[Tuple[int, str]],
    sequence: int,
) -> Chunk:
    doc_id = _doc_id(path)
    start, end = rows[0][0], rows[-1][0]
    chunk_id = f"{doc_id}:p{start}-p{end}:c{sequence:04d}"
    body = "\n".join(text for _, text in rows)
    text = f"Section: {section}\n{body}" if section else body
    return Chunk(
        chunk_id=chunk_id,
        text=text,
        doc_id=doc_id,
        source_file=path.name,
        section=section or "Document introduction",
        paragraph_start=start,
        paragraph_end=end,
        topic=_topic(section),
    )


def chunk_booklet(
    path: Path,
    target_chars: int = 1_800,
    max_chars: int = 2_600,
    overlap_paragraphs: int = 1,
) -> List[Chunk]:
    """Chunk one workbook without crossing detected section boundaries."""
    if target_chars < 1 or target_chars > max_chars:
        raise ValueError("target_chars must be positive and <= max_chars")
    if overlap_paragraphs < 0:
        raise ValueError("overlap_paragraphs must be non-negative")
    rows = load_booklet_rows(path)
    chunks: List[Chunk] = []
    section = "Document introduction"
    pending: List[Tuple[int, str]] = []
    pending_chars = 0
    sequence = 1

    def flush(retain_overlap: bool = True) -> None:
        nonlocal pending, pending_chars, sequence
        if pending:
            chunks.append(_emit_chunk(path, section, pending, sequence))
            sequence += 1
            # Preserve one paragraph of local context between long-section chunks.
            pending = (
                pending[-overlap_paragraphs:]
                if retain_overlap and overlap_paragraphs and pending_chars >= target_chars
                else []
            )
            pending_chars = sum(len(text) for _, text in pending)

    body_limit = max_chars - 200
    if body_limit < 1:
        raise ValueError("max_chars must leave room for section metadata")
    expanded_rows = [
        (paragraph_number, part)
        for paragraph_number, text in rows
        for part in split_long_text(text, body_limit)
    ]

    for paragraph_number, text in expanded_rows:
        if is_heading(text):
            # Consecutive headings are common in contents pages. Keep the most
            # specific heading instead of producing a heading-only chunk.
            if not (len(pending) == 1 and is_heading(pending[0][1])):
                flush()
            pending = []
            pending_chars = 0
            section = clean_heading(text)
            pending.append((paragraph_number, text))
            pending_chars = len(text)
            continue

        projected = pending_chars + len(text) + 1
        if pending and projected > max_chars:
            flush()
        pending.append((paragraph_number, text))
        pending_chars += len(text) + 1
        if pending_chars >= target_chars and text.endswith((".", ":", ";")):
            flush()

    flush(retain_overlap=False)
    return chunks


def chunk_documents(
    paths: Iterable[Path],
    target_chars: int = 1_800,
    max_chars: int = 2_600,
    overlap_paragraphs: int = 1,
) -> List[Chunk]:
    chunks: List[Chunk] = []
    for path in sorted(Path(item) for item in paths):
        chunks.extend(chunk_booklet(path, target_chars, max_chars, overlap_paragraphs))
    return chunks


def write_chunks(chunks: Iterable[Chunk], path: Path) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with path.open("w", encoding="utf-8") as handle:
        for chunk in chunks:
            handle.write(json.dumps(chunk.to_dict(), ensure_ascii=False) + "\n")
            count += 1
    return count


def read_chunks(path: Path) -> List[Chunk]:
    with path.open(encoding="utf-8") as handle:
        return [Chunk.from_dict(json.loads(line)) for line in handle if line.strip()]
