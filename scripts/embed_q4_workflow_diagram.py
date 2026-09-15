"""Embed the generated Q4 workflow diagram in the submission DOCX."""

from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "shared_output/Q4_Graph_Multi_Agent/graph-based-multi-agent.docx"
DIAGRAM = ROOT / "shared_output/Q4_Graph_Multi_Agent/Q4_GRAPH_MULTI_AGENT_WORKFLOW.png"


def main():
    document = Document(DOCX)
    for paragraph in list(document.paragraphs):
        if "<w:drawing" in paragraph._p.xml:
            paragraph._element.getparent().remove(paragraph._element)

    for paragraph in list(document.paragraphs):
        if paragraph.text.startswith("The workflow is a bounded state graph"):
            paragraph._element.getparent().remove(paragraph._element)

    workflow_heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "1. Graph workflow design and topology"
    )
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = 1
    image_paragraph.add_run().add_picture(str(DIAGRAM), width=Inches(6.95))
    workflow_heading._p.addnext(image_paragraph._p)
    document.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
