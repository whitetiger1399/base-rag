"""Embed the generated Q3 workflow diagram in the submission DOCX."""

from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "shared_output/Q3_Guardrails_Eval_Agents/Q3_Guardrails_Eval_Agents.docx"
DIAGRAM = ROOT / "shared_output/Q3_Guardrails_Eval_Agents/Q3_GUARDRAIL_EVAL_AGENT_WORKFLOW.png"


def insert_after(paragraph, new_paragraph):
    paragraph._p.addnext(new_paragraph._p)


def main():
    document = Document(DOCX)

    # Make the operation repeatable by removing only paragraphs containing a drawing.
    for paragraph in list(document.paragraphs):
        if "<w:drawing" in paragraph._p.xml:
            paragraph._element.getparent().remove(paragraph._element)

    workflow = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Workflow. START")
    )
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = 1
    image_paragraph.add_run().add_picture(str(DIAGRAM), width=Inches(6.95))
    insert_after(workflow, image_paragraph)

    document.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
