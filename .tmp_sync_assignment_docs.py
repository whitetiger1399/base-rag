from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path("shared_output")


def find_paragraph(document, prefix):
    return next(p for p in document.paragraphs if p.text.startswith(prefix))


def insert_after(paragraph, text, style="Normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    created = Paragraph(new_p, paragraph._parent)
    created.style = style
    created.add_run(text)
    return created


# Q1: connect the retrieval design to the measured three-question evidence.
q1_path = ROOT / "Q1_Retrieval_Design/Q1_Retrieval_Design.docx"
q1 = Document(q1_path)
marker = "Measured validation sample."
if not any(p.text.startswith(marker) for p in q1.paragraphs):
    anchor = find_paragraph(q1, "Choosing top_k.")
    insert_after(
        anchor,
        marker
        + " In the three-question local Ragas sample (Q220, Q1016, and Q1226), "
        "context precision and context recall were 1.0000 for every question. This "
        "supports the current hybrid retrieval configuration for those examples, "
        "but it is not a substitute for deterministic Recall@k/MRR on the frozen "
        "golden set or evidence that hybrid retrieval always outperforms vector-only search.",
    )
q1.save(q1_path)


# Q3: synchronize the guardrail narrative with the hardened reference snippet.
q3_path = ROOT / "Q3_Guardrails_Eval_Agents/Q3_Guardrails_Eval_Agents.docx"
q3 = Document(q3_path)
find_paragraph(q3, "Defense in depth.").text = (
    "Defense in depth. Normalize untrusted query and chunk text with Unicode "
    "compatibility folding, invisible-character removal, and conservative "
    "homoglyph/leet mapping. Check word-preserving and compact forms with "
    "high-signal patterns, then use an independently tested local semantic "
    "detector for injection intent beyond regex. Detector errors fail closed. "
    "Quarantine flagged chunks by ID, verify only the safe set, and emit structured "
    "audit events without logging malicious source text."
)
find_paragraph(q3, "Run retrieval first").text = (
    "Run retrieval first and compute Recall@k and MRR from annotated chunk IDs. "
    "Then generate answers and parse citations. Reject unknown IDs deterministically. "
    "Split the answer into atomic factual claims and ask qwen3:8b, at temperature 0 "
    "with JSON output, for total claims, supported claims, and unsupported claims. "
    "Compute faithfulness deterministically as supported claims divided by total "
    "claims, allowing fractional credit. A second human reviewer audits disagreements, "
    "all safety failures, and a random sample."
)
find_paragraph(q3, "Query → Retriever").text = (
    "Query → normalized Safety screen → Retriever → chunk Safety screen → optional "
    "embedding-similarity gate → Verifier. If sufficient, the Answerer receives only "
    "approved chunks and no tool handles, then the system validates answer type, "
    "citation IDs, and fractional claim support. Unsafe chunks are quarantined and "
    "the reduced set is rechecked. Permit at most two retrieval retries; empty safe "
    "evidence, no-progress refinement, invalid output, or faithfulness below 0.90 "
    "routes to the exact abstention response."
)
find_paragraph(q3, "The companion snippets_q3.py").text = (
    "The companion snippets_q3.py implements security normalization, layered regex "
    "and optional semantic injection detection, a replaceable non-LLM evidence "
    "similarity check, structured security audit callbacks, fractional local-Ollama "
    "faithfulness scoring, and dependency-injected orchestration. The Answerer is "
    "passed only verifier-approved evidence and exposes no tool interface through "
    "this contract."
)
guardrail_table = q3.tables[0]
guardrail_table.rows[1].cells[1].text = (
    "Normalize Unicode/spacing tricks; run word and compact patterns plus optional "
    "semantic intent detection."
)
guardrail_table.rows[1].cells[2].text = (
    "Quarantine by ID, log detection layer, and re-verify safe evidence; fail closed "
    "if the semantic detector errors."
)
q3.save(q3_path)


# Q4: propagate the same controls into the graph's Safety and validation nodes.
q4_path = ROOT / "Q4_Graph_Multi_Agent/graph-based-multi-agent.docx"
q4 = Document(q4_path)
find_paragraph(q4, "The graph turns").text = (
    "The graph turns the Q3 agent contracts into deterministic control flow. Nodes "
    "perform one responsibility and return structured state; conditional edges decide "
    "whether to answer, refine retrieval, filter unsafe evidence, or abstain. Local "
    "qwen3:8b can implement Verifier and Answerer, while Unicode normalization, "
    "layered injection detection, routing, retry limits, citation allow-lists, and "
    "audit events remain application controlled."
)
find_paragraph(q4, "Retrieved instructions are data").text = (
    "Retrieved instructions are data, never commands. The Safety node applies "
    "security-only Unicode normalization, removes invisible formatting, maps a "
    "conservative set of homoglyph/leet variants, and checks both word-preserving "
    "and compact forms. A replaceable semantic detector handles intent outside fixed "
    "patterns and fails closed on detector errors. Flagged chunks are quarantined by "
    "ID before verification; raw malicious text is not written to audit logs."
)
find_paragraph(q4, "The Answerer sees only approved chunks").text = (
    "The Answerer sees only approved chunks and the citation allow-list; the graph "
    "does not provide it with tool handles. A deterministic post-check requires a "
    "plain-text answer and at least one allow-listed [chunk_id] for factual output. "
    "The Q3 claim checker returns supported claims divided by total factual claims, "
    "so partial support receives a fractional score. A score below 0.90 routes to "
    "Abstain or one controlled revision, never an open-ended loop."
)
find_paragraph(q4, "Timeouts, invalid schemas").text = (
    "Timeouts, classifier failures, invalid schemas, missing indexes, and empty "
    "results are represented as explicit state outcomes. Log node duration, attempt, "
    "retrieved/safe/approved IDs, quarantine detection layer, evidence-similarity "
    "score, route, verdict, citation result, fractional faithfulness, and terminal "
    "reason. Redact raw query or passage text when it may contain sensitive or "
    "malicious information."
)
node_table = q4.tables[2]
node_table.rows[2].cells[1].text = (
    "Normalize obfuscation; run word/compact patterns and optional semantic intent "
    "detection; quarantine by ID and emit safe audit metadata."
)
node_table.rows[2].cells[2].text = (
    "safe_chunks and quarantined_ids are disjoint; detector errors fail closed."
)
q4.save(q4_path)
